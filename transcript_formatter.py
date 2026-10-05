"""Helpers for formatting speaker-labeled transcripts."""
from io import BytesIO
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def format_transcript_by_speaker(text: str, speaker_map: dict[str, str]) -> str:
    """Replace bracketed speaker labels with readable ``Name:`` labels."""
    def replace_label(match: re.Match[str]) -> str:
        speaker_id = match.group(1)
        name = speaker_map.get(speaker_id, "").strip()
        return f"{name or speaker_id}:"

    return re.sub(r"\[([^\]\r\n]+)\]", replace_label, text)


def transcript_to_docx(text: str) -> bytes:
    """Create a DOCX document with speaker names emphasized."""
    document = Document()
    document.core_properties.language = "pt-BR"

    normal_style = document.styles["Normal"]
    style_properties = normal_style._element.get_or_add_rPr()
    language = style_properties.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        style_properties.append(language)
    language.set(qn("w:val"), "pt-BR")

    for block in text.split("\n\n"):
        if not block.strip():
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        speaker, separator, utterance = block.partition(": ")
        if separator:
            paragraph.add_run(f"{speaker}: ").bold = True
            paragraph.add_run(utterance)
        else:
            paragraph.add_run(block)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
