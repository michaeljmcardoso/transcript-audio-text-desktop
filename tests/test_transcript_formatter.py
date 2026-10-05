import unittest
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

from transcript_formatter import format_transcript_by_speaker, transcript_to_docx


class TranscriptFormatterTests(unittest.TestCase):
    def test_formats_named_speakers_without_brackets(self):
        transcript = (
            "[SPEAKER_00] Hoje é 12 de junho.\n\n"
            "[SPEAKER_01] Eu sou Lizonete."
        )

        formatted = format_transcript_by_speaker(
            transcript,
            {"SPEAKER_00": "Entrevistador", "SPEAKER_01": "Lizonete"},
        )

        self.assertEqual(
            formatted,
            "Entrevistador: Hoje é 12 de junho.\n\nLizonete: Eu sou Lizonete.",
        )
        self.assertNotIn("[", formatted)

    def test_uses_speaker_id_when_name_is_missing(self):
        transcript = "[SPEAKER_00] Olá.\n\n[DESCONHECIDO] Áudio inaudível."

        self.assertEqual(
            format_transcript_by_speaker(transcript, {"SPEAKER_00": "  "}),
            "SPEAKER_00: Olá.\n\nDESCONHECIDO: Áudio inaudível.",
        )

    def test_docx_export_preserves_paragraphs_and_emphasizes_speaker_names(self):
        transcript = "Entrevistador: Bom dia.\n\nLizonete: Eu moro aqui."

        exported = transcript_to_docx(transcript)
        document = Document(BytesIO(exported))

        self.assertEqual(
            [paragraph.text for paragraph in document.paragraphs],
            ["Entrevistador: Bom dia.", "Lizonete: Eu moro aqui."],
        )
        self.assertTrue(document.paragraphs[0].runs[0].bold)
        self.assertFalse(document.paragraphs[0].runs[1].bold)
        self.assertEqual(document.core_properties.language, "pt-BR")
        self.assertEqual(
            document.styles["Normal"]._element.rPr.find(qn("w:lang")).get(qn("w:val")),
            "pt-BR",
        )


if __name__ == "__main__":
    unittest.main()
